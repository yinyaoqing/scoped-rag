"""Files in, plain text out.

A batch keeps going when one file fails: ingesting a folder of real documents and
having the whole run die on a single corrupt PDF is worse than skipping it and
saying so at the end.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv"}


class UnsupportedFormat(RuntimeError):
    pass


def parse_file(path: str | Path) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return p.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        return _parse_pdf(p)
    if suffix == ".docx":
        return _parse_docx(p)
    raise UnsupportedFormat(f"{p.name}: no parser for {suffix or 'files without a suffix'}")


def _parse_pdf(p: Path) -> str:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError as exc:  # pragma: no cover - optional extra
        raise UnsupportedFormat("PDF needs the 'parse' extra: uv sync --extra parse") from exc
    reader = PdfReader(str(p))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _parse_docx(p: Path) -> str:
    try:
        import docx
    except ModuleNotFoundError as exc:  # pragma: no cover - optional extra
        raise UnsupportedFormat("DOCX needs the 'parse' extra: uv sync --extra parse") from exc
    document = docx.Document(str(p))
    parts = [para.text.strip() for para in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(part for part in parts if part)


@dataclass(frozen=True)
class Parsed:
    path: Path
    text: str


@dataclass(frozen=True)
class Failed:
    path: Path
    reason: str


def parse_many(paths: list[str | Path]) -> tuple[list[Parsed], list[Failed]]:
    ok: list[Parsed] = []
    bad: list[Failed] = []
    for path in paths:
        p = Path(path)
        try:
            text = parse_file(p)
        except Exception as exc:  # noqa: BLE001 - one bad file must not stop the batch
            bad.append(Failed(p, str(exc)))
            continue
        if text.strip():
            ok.append(Parsed(p, text))
        else:
            bad.append(Failed(p, "parsed to no text"))
    return ok, bad
